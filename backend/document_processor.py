import io
import re
from typing import List, Dict, Any, Optional, Tuple
from docx import Document
from docx.text.paragraph import Paragraph

class DocumentBlock:
    def __init__(self, block_id: str, text: str, location_type: str, metadata: Dict[str, Any]):
        self.block_id = block_id
        self.text = text
        self.location_type = location_type  # "body" or "table"
        self.metadata = metadata

class DocxProcessor:
    """
    Handles extracting text blocks (paragraphs & tables) from .docx files,
    tracking their locations, and applying approved textual corrections
    while preserving document formatting and structure.
    """

    @staticmethod
    def extract_blocks(file_bytes: bytes) -> Tuple[Document, List[Dict[str, Any]]]:
        """
        Extracts all textual blocks from docx bytes.
        Returns the docx Document object and a list of block metadata dictionaries.
        """
        doc = Document(io.BytesIO(file_bytes))
        blocks: List[Dict[str, Any]] = []

        # 1. Body paragraphs
        for p_idx, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if text:
                style_name = p.style.name if p.style else "Normal"
                is_title = bool(
                    "title" in style_name.lower() or 
                    "heading" in style_name.lower() or 
                    text.lower().startswith("title:") or
                    text.lower().startswith("title -") or
                    text.lower().startswith("title")
                )
                blocks.append({
                    "id": f"p_{p_idx}",
                    "text": p.text,
                    "location_type": "body",
                    "p_index": p_idx,
                    "style": style_name,
                    "is_title_or_heading": is_title
                })

        # 2. Table paragraphs
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    for cp_idx, p in enumerate(cell.paragraphs):
                        text = p.text.strip()
                        if text:
                            style_name = p.style.name if p.style else "Normal"
                            is_title = bool(
                                "title" in style_name.lower() or 
                                "heading" in style_name.lower() or 
                                text.lower().startswith("title:") or
                                (r_idx == 0 and len(table.rows) > 1) # header row
                            )
                            blocks.append({
                                "id": f"tbl_{t_idx}_{r_idx}_{c_idx}_{cp_idx}",
                                "text": p.text,
                                "location_type": "table",
                                "table_index": t_idx,
                                "row_index": r_idx,
                                "col_index": c_idx,
                                "p_index": cp_idx,
                                "style": style_name,
                                "is_title_or_heading": is_title
                            })

        return doc, blocks

    @staticmethod
    def _replace_in_paragraph(paragraph: Paragraph, old_text: str, new_text: str) -> bool:
        """
        Replaces old_text with new_text in the given paragraph while preserving run-level formatting.
        Returns True if replacement was performed.
        """
        if not old_text or old_text not in paragraph.text:
            return False

        # Try simple single-run replacement first (best formatting preservation)
        for run in paragraph.runs:
            if old_text in run.text:
                run.text = run.text.replace(old_text, new_text, 1)
                return True

        # If old_text spans across multiple runs:
        # We can map character positions across runs and replace cleanly
        runs = paragraph.runs
        if not runs:
            paragraph.text = paragraph.text.replace(old_text, new_text, 1)
            return True

        full_text = "".join(r.text for r in runs)
        start_pos = full_text.find(old_text)
        if start_pos == -1:
            return False

        end_pos = start_pos + len(old_text)

        # Find which runs intersect [start_pos, end_pos]
        cur_pos = 0
        run_spans = []
        for run in runs:
            run_len = len(run.text)
            run_spans.append((cur_pos, cur_pos + run_len, run))
            cur_pos += run_len

        # Distribute the replacement:
        # Put the new_text into the first intersecting run, and clear the matched segment in subsequent runs
        first_run_updated = False
        for r_start, r_end, run in run_spans:
            if r_end <= start_pos or r_start >= end_pos:
                # Outside the target range
                continue

            # This run intersects [start_pos, end_pos]
            overlap_start = max(0, start_pos - r_start)
            overlap_end = min(len(run.text), end_pos - r_start)

            prefix = run.text[:overlap_start]
            suffix = run.text[overlap_end:]

            if not first_run_updated:
                run.text = prefix + new_text + suffix
                first_run_updated = True
            else:
                run.text = prefix + suffix

        return True

    @classmethod
    def apply_corrections(
        cls,
        doc: Document,
        corrections: List[Dict[str, Any]]
    ) -> bytes:
        """
        Applies a list of approved corrections to the Document object.
        Each correction has:
          - block_id: str
          - original_text: str
          - replacement_text: str

        Returns the saved .docx as bytes.
        """
        # Group corrections by block_id
        corrections_by_block: Dict[str, List[Dict[str, Any]]] = {}
        for corr in corrections:
            b_id = corr.get("block_id")
            if b_id:
                corrections_by_block.setdefault(b_id, []).append(corr)

        # 1. Apply to body paragraphs
        for p_idx, p in enumerate(doc.paragraphs):
            b_id = f"p_{p_idx}"
            if b_id in corrections_by_block:
                for corr in corrections_by_block[b_id]:
                    orig = corr.get("original_text", "")
                    repl = corr.get("replacement_text", "")
                    if orig and repl is not None:
                        cls._replace_in_paragraph(p, orig, repl)

        # 2. Apply to table paragraphs
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    for cp_idx, p in enumerate(cell.paragraphs):
                        b_id = f"tbl_{t_idx}_{r_idx}_{c_idx}_{cp_idx}"
                        if b_id in corrections_by_block:
                            for corr in corrections_by_block[b_id]:
                                orig = corr.get("original_text", "")
                                repl = corr.get("replacement_text", "")
                                if orig and repl is not None:
                                    cls._replace_in_paragraph(p, orig, repl)

        out_stream = io.BytesIO()
        doc.save(out_stream)
        out_stream.seek(0)
        return out_stream.getvalue()

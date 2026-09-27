import io
import unittest
from docx import Document
from backend.document_processor import DocxProcessor

class TestDocxProcessor(unittest.TestCase):
    def test_block_extraction_and_replacement(self):
        # Create test document
        doc = Document()
        p1 = doc.add_paragraph()
        run1 = p1.add_run("This is an ")
        run2 = p1.add_run("erronous")
        run2.bold = True
        run3 = p1.add_run(" statemnt in the report.")

        # Save to bytes
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        file_bytes = buf.getvalue()

        # Extract blocks
        loaded_doc, blocks = DocxProcessor.extract_blocks(file_bytes)
        self.assertEqual(len(blocks), 1)
        self.assertEqual(blocks[0]["id"], "p_0")
        self.assertIn("erronous", blocks[0]["text"])

        # Apply corrections
        corrections = [
            {
                "block_id": "p_0",
                "original_text": "erronous",
                "replacement_text": "erroneous"
            },
            {
                "block_id": "p_0",
                "original_text": "statemnt",
                "replacement_text": "statement"
            }
        ]

        output_bytes = DocxProcessor.apply_corrections(loaded_doc, corrections)
        self.assertIsNotNone(output_bytes)

        # Verify applied corrections in new document
        res_doc = Document(io.BytesIO(output_bytes))
        self.assertIn("erroneous", res_doc.paragraphs[0].text)
        self.assertIn("statement", res_doc.paragraphs[0].text)
        self.assertNotIn("erronous", res_doc.paragraphs[0].text)

        # Verify bold run was preserved
        bold_runs = [r for r in res_doc.paragraphs[0].runs if r.bold]
        self.assertTrue(len(bold_runs) > 0)
        self.assertEqual(bold_runs[0].text, "erroneous")
        print("[SUCCESS] All docx replacement and formatting tests passed!")

if __name__ == "__main__":
    unittest.main()
